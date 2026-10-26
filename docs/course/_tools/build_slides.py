"""Build the progress presentation (Tasks 1–3) as PPTX.
Needs python-pptx:  python -m pip install python-pptx
Run:  python docs/course/_tools/build_slides.py
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
COURSE = os.path.dirname(HERE)
T1 = os.path.join(COURSE, "1. Software Process", "diagrams")
T2 = os.path.join(COURSE, "2. Requirements", "diagrams")
T3 = os.path.join(COURSE, "3. Design", "diagrams")
OUT = os.path.join(COURSE, "slides")
os.makedirs(OUT, exist_ok=True)

FONT = "Arial"
INK = RGBColor(0x1F, 0x29, 0x33)
MUTED = RGBColor(0x52, 0x60, 0x6D)
BLUE = RGBColor(0x25, 0x63, 0xEB)
GREEN = RGBColor(0x15, 0x80, 0x3D)
AMBER = RGBColor(0xB4, 0x53, 0x09)
LIGHT = RGBColor(0xF5, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
TOTAL = []


def tb(slide, x, y, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, color
        p.alignment = align
    return box


def bullets(slide, x, y, w, h, items, size=18, gap=8):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        sub = it.startswith("  ")
        head, _, rest = it.strip().partition("|")
        r = p.add_run()
        r.text = ("–  " if sub else "•  ") + head
        r.font.name, r.font.size = FONT, Pt(size - (3 if sub else 0))
        r.font.bold = bool(rest)
        r.font.color.rgb = MUTED if sub else INK
        if rest:
            r2 = p.add_run()
            r2.text = rest
            r2.font.name, r2.font.size, r2.font.color.rgb = FONT, Pt(size), INK
        p.space_after = Pt(gap)
        if sub:
            p.level = 1
    return box


def slide(title, kicker=None):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), SH)
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    if kicker:
        tb(s, Inches(0.55), Inches(0.28), Inches(12), Inches(0.4), kicker.upper(), size=12, bold=True, color=BLUE)
    tb(s, Inches(0.55), Inches(0.55), Inches(12.3), Inches(0.8), title, size=28, bold=True)
    TOTAL.append(s)
    return s


def image(s, path, x, y, maxw, maxh):
    w, h = Image.open(path).size
    scale = min(maxw / w, maxh / h)
    pw, ph = int(w * scale), int(h * scale)
    s.shapes.add_picture(path, x + (maxw - pw) // 2, y + (maxh - ph) // 2, pw, ph)


def table(s, x, y, w, rows, col_w, size=13, header_fill=BLUE, row_h=0.42):
    shape = s.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(row_h * len(rows)))
    t = shape.table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = str(val)
            p = c.text_frame.paragraphs[0]
            for r in p.runs:
                r.font.name, r.font.size = FONT, Pt(size)
                r.font.bold = i == 0
                r.font.color.rgb = WHITE if i == 0 else INK
            c.fill.solid()
            c.fill.fore_color.rgb = header_fill if i == 0 else (LIGHT if i % 2 else WHITE)
    return t


def footer_numbers():
    n = len(TOTAL)
    for i, s in enumerate(TOTAL, 1):
        tb(s, SW - Inches(2.6), SH - Inches(0.45), Inches(2.3), Inches(0.3),
           f"Proxy Courses · {i}/{n}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


BODY_X, BODY_Y = Inches(0.55), Inches(1.45)
BODY_W, BODY_H = SW - Inches(1.1), SH - Inches(1.95)

# 1 ── Title
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
bg.fill.solid(); bg.fill.fore_color.rgb = BLUE; bg.line.fill.background()
tb(s, Inches(0.9), Inches(1.4), Inches(11.5), Inches(1.2), "Proxy Courses", size=54, bold=True, color=WHITE)
tb(s, Inches(0.9), Inches(2.45), Inches(11.5), Inches(0.8),
   "A peer-tutoring and advising platform for IUB students", size=24, color=WHITE)
tb(s, Inches(0.9), Inches(3.25), Inches(11.5), Inches(0.6),
   "Progress review · Tasks 1–3: Software Process, Requirements, Design", size=20, bold=True, color=WHITE)
tb(s, Inches(0.9), Inches(4.4), Inches(6.5), Inches(2.2), [
    "Syed Mehedi Hussain · 2310690", "Arafat Hossain · 2420834",
    "Rukaiya Tabassum · 2420183", "Meharab Hossen Opi · 2420832"], size=18, color=WHITE)
tb(s, Inches(7.6), Inches(4.4), Inches(5.2), Inches(2.2), [
    "CSE451 / CSC445 Software Engineering", "Autumn 2026 · Independent University, Bangladesh",
    "Faculty: Md Masum Billah"], size=16, color=WHITE)
TOTAL.append(s)

# 2 ── Problem and idea
s = slide("The problem and our answer", "Context")
bullets(s, BODY_X, BODY_Y, Inches(6.1), BODY_H, [
    "Problem|",
    "  Learners struggle in specific IUB courses; coaching is costly and off-syllabus",
    "  Seniors who aced the course are found only by luck (Facebook groups)",
    "  No trust: no grade check, no reviews, no refunds",
    "  Faculty and club advice is scattered",
], size=20)
bullets(s, Inches(6.9), BODY_Y, Inches(5.9), BODY_H, [
    "Proxy Courses|",
    "  Tutors with an admin-verified A / A- in the course",
    "  ৳500 semester subscription; 70/30 split (80/20 for top tutors)",
    "  Free intro video + free live demo before paying; clear refund rule",
    "  Free booking with faculty advisors and club executives",
    "  Responsive web app first; one API ready for mobile later",
], size=20)

# 3 ── Status
s = slide("Where we are", "Status")
table(s, BODY_X, BODY_Y, Inches(8.2), [
    ["Task", "Marks", "Status"],
    ["1  Software Process", "10", "Done: 3 / 3 items, 1 diagram"],
    ["2  Requirements Engineering", "22", "Done: 16 / 16 items, 6 diagrams"],
    ["3  Software Design", "16", "Done: 9 / 9 items, 6 diagrams"],
    ["4  Product Metrics", "14", "Next"],
    ["5  Quality Attributes", "10", "Planned"],
    ["6  Testing", "18", "Planned (test plan drafted)"],
    ["7  Project Management", "18", "Planned"],
    ["8  Risk Management", "12", "Planned (6 risks drafted)"],
], [3.6, 1.0, 3.6], size=15, row_h=0.5)
bullets(s, Inches(9.1), BODY_Y, Inches(3.8), BODY_H, [
    "Also written|",
    "  PRD: 36 FRs, 12 NFRs",
    "  Architecture, data model, API",
    "  Roadmap M1–M4",
    "  Stack ADR: Next.js + PostgreSQL",
    "  Payments ADR: mock gateway → SSLCommerz",
    "48 of 120 marks' worth of work drafted|",
], size=17)

# 4 ── Task 1 process model
s = slide("Task 1 · Choosing the process model", "Software process")
table(s, BODY_X, BODY_Y, Inches(8.4), [
    ["Criterion", "Weight", "Waterfall", "Spiral", "Scrum + XP"],
    ["C1 Copes with changing requirements", "0.25", "1", "4", "5"],
    ["C2 Early working software", "0.20", "1", "3", "5"],
    ["C3 Fits 14 weeks, task every 2 weeks", "0.20", "3", "2", "5"],
    ["C4 Produces graded documents", "0.20", "5", "4", "3"],
    ["C5 Risk control, small team", "0.15", "2", "5", "4"],
    ["Weighted score", "1.00", "2.35", "3.55", "4.45"],
], [3.9, 1.0, 1.2, 1.1, 1.2], size=14, row_h=0.52)
bullets(s, Inches(9.3), BODY_Y, Inches(3.7), BODY_H, [
    "Chosen: Scrum + XP + MoSCoW|",
    "  Pricing and refund rules will change after interviews",
    "  7 two-week sprints = 7 task deadlines",
    "  XP: pair work, PR review, CI on every push",
    "  Weak spot (docs) fixed: each sprint's Definition of Done includes its task document",
], size=17)

# 5 ── Task 1 lifecycle
s = slide("Task 1 · Tailored life cycle", "Software process")
image(s, os.path.join(T1, "lifecycle.png"), BODY_X, BODY_Y, BODY_W, BODY_H)

# 6 ── Stakeholders
s = slide("Task 2 · 12 stakeholders on a power–interest grid", "Requirements")
image(s, os.path.join(T2, "power-interest-grid.png"), BODY_X, BODY_Y, Inches(7.2), BODY_H)
bullets(s, Inches(8.0), BODY_Y, Inches(4.9), BODY_H, [
    "Manage closely|: learners, tutors, course faculty, team",
    "Keep satisfied|: IUB administration, payment gateway, faculty of tutored courses",
    "Keep informed|: advisors, club executives, parents",
    "Monitor|: coaching centres, IUB IT",
    "Elicitation|: 9 interviews · survey (≥ 50) · analysis of 30 Facebook posts",
], size=17)

# 7 ── Requirements catalogue
s = slide("Task 2 · Requirements catalogue and user stories", "Requirements")
table(s, BODY_X, BODY_Y, Inches(5.8), [
    ["Catalogue", "Count"],
    ["Functional requirements", "36"],
    ["   Must / Should / Could", "28 / 7 / 1"],
    ["   Won't (v1)", "6 items"],
    ["Non-functional, each with a metric", "12"],
    ["User stories (INVEST-checked)", "6"],
    ["Use cases", "15"],
], [4.3, 1.5], size=15, row_h=0.5)
bullets(s, Inches(6.8), BODY_Y, Inches(6.1), BODY_H, [
    "Example NFR metrics|",
    "  Search p95 ≤ 500 ms with 1,000 offerings",
    "  ≥ 99.0% monthly uptime (≤ 7.3 h down)",
    "  0 transcripts reachable by other users",
    "  WCAG 2.1 AA: 0 serious axe-core issues",
    "INVEST caught 2 problems|",
    "  US-3 too big → split into mock-pay and real gateway",
    "  US-5 depended on US-4 → seeded test data",
], size=17)

# 8 ── Use case diagram
s = slide("Task 2 · Use-case diagram (15 use cases, 6 actors)", "Requirements")
image(s, os.path.join(T2, "use-case-diagram.png"), BODY_X, BODY_Y, Inches(8.6), BODY_H)
bullets(s, Inches(9.35), BODY_Y, Inches(3.6), BODY_H, [
    "«include»|: Subscribe always includes Pay via gateway",
    "«extend»|: Book free demo extends View offering [seats free]",
    "Generalisation|: a tutor is also a learner",
    "Fully dressed|: UC-05 Subscribe, with 6 extensions (failed payment, forged callback, duplicate callback…)",
], size=15)

# 9 ── ER
s = slide("Task 2 · ER diagram (crow's foot)", "Requirements")
image(s, os.path.join(T2, "er-diagram.png"), BODY_X, BODY_Y, Inches(9.0), BODY_H)
bullets(s, Inches(9.75), BODY_Y, Inches(3.2), BODY_H, [
    "13 entities|, PKs underlined",
    "Weak entity|: MATERIAL (owner OFFERING + material_no)",
    "Associative|: DEMO_BOOKING",
    "Rules in the model|: REVIEW and REFUND hang off SUBSCRIPTION (0..1), so only subscribers can review",
], size=15)

# 10 ── State
s = slide("Task 2 · Subscription state diagram", "Requirements")
image(s, os.path.join(T2, "state-diagram.png"), BODY_X, BODY_Y, Inches(7.0), BODY_H)
bullets(s, Inches(7.9), BODY_Y, Inches(5.0), BODY_H, [
    "6 states|: PendingPayment, Active, Cancelled, RefundRequested, Refunded, Expired",
    "Every transition = event [guard] / action|",
    "  paymentSucceeded [signature valid ∧ amount = ৳500 ∧ txn new] / mark PAID, store split",
    "  requestRefund [inside window] / add to admin queue",
    "Guards prevent double charging and access without payment|",
], size=16)

# 11 ── DFD
s = slide("Task 2 · Level-1 DFD: “Subscribe to an offering”", "Requirements")
image(s, os.path.join(T2, "dfd-level1.png"), BODY_X, BODY_Y, BODY_W, Inches(4.4))
tb(s, BODY_X, Inches(6.15), BODY_W, Inches(0.8),
   "5 processes (find → order → pay → activate & split → notify) · 4 external entities · 5 data stores · "
   "flows balance with the context-level process", size=16, color=MUTED)

# 12 ── Win-win, review, traceability
s = slide("Task 2 · Negotiation, review and traceability", "Requirements")
bullets(s, BODY_X, BODY_Y, Inches(6.2), BODY_H, [
    "Win-win 1: refunds vs tutor income|",
    "  Full refund within 14 days and before the 3rd live class; later only by dispute (to confirm with interviews)",
    "Win-win 2: trust vs transcript privacy|",
    "  Public sees “✓ Verified A in CSE211”; only admins see the transcript",
    "8-question review on FR-16, FR-25, NFR-01|",
    "  Found 3 ambiguities, 2 gaps, 1 conflict, all fixed",
], size=16)
table(s, Inches(7.0), BODY_Y, Inches(5.8), [
    ["Req.", "Use case", "Module", "Test"],
    ["FR-16", "UC-02", "M4 Discovery", "TC-06"],
    ["FR-19", "UC-04", "M5 Booking", "TC-07"],
    ["FR-22", "UC-06", "M7 Payments", "TC-09"],
    ["FR-25", "UC-08", "M6 Subscriptions", "TC-12"],
    ["NFR-05", "UC-10/12", "M2, M12", "TC-18"],
    ["… 21 rows in total", "", "", ""],
], [1.1, 1.3, 2.2, 1.2], size=13, row_h=0.45)

# 13 ── Change control
s = slide("Task 2 · Change control and the requirements baseline", "Requirements")
image(s, os.path.join(T2, "change-control.png"), BODY_X, BODY_Y, Inches(6.3), BODY_H)
bullets(s, Inches(7.2), BODY_Y, Inches(5.7), BODY_H, [
    "Baseline v1.0|: PRD + Task 2 models, approved at the Sprint 2 review",
    "  Frozen as git tag req-baseline-v1.0",
    "After baselining|: change only through a CR",
    "  Impact analysis uses the traceability matrix",
    "  CCB: product owner + scrum master + QA lead",
    "  Approved → v1.1 with a new tag",
    "Scope or graded-deliverable changes are checked with the faculty first|",
], size=16)

# 14 ── Modules
s = slide("Task 3 · 12 modules in three layers", "Design")
image(s, os.path.join(T3, "module-diagram.png"), BODY_X, BODY_Y, Inches(7.0), BODY_H)
bullets(s, Inches(7.8), BODY_Y, Inches(5.1), BODY_H, [
    "Each module|: responsibility, interface, hidden information, dependencies",
    "No cycles|: M6 Subscriptions → M7 Payments, never back",
    "Hidden|: only M7 knows which gateway is used",
    "Cross-cutting|: roles, audit log, validation, error mapping",
], size=17)

# 15 ── God class
s = slide("Task 3 · Fixing a god class", "Design")
image(s, os.path.join(T3, "refactor-before.png"), BODY_X, BODY_Y, Inches(6.1), Inches(3.0))
image(s, os.path.join(T3, "refactor-after.png"), Inches(6.75), BODY_Y, Inches(6.1), Inches(3.0))
bullets(s, BODY_X, Inches(4.7), BODY_W, Inches(2.3), [
    "Before|: SubscriptionManager had coincidental and logical cohesion, and content, common, control and external coupling (15 methods, 9 collaborators)",
    "After|: functional cohesion, data coupling and a PaymentGateway interface (mock or SSLCommerz). The control flag is replaced by polymorphism.",
], size=16)

# 16 ── Golden rules + personas
s = slide("Task 3 · Mandel's golden rules, personas, task analysis", "Design")
table(s, BODY_X, BODY_Y, Inches(6.3), [
    ["Golden rule", "In Proxy Courses"],
    ["User in control", "✕ on every checkout step; cancel demo; withdraw refund"],
    ["Reduce memory load", "Step 2 of 3; refund as a date; last method pre-selected"],
    ["Consistent interface", "One blue bottom button; same status colours; same words"],
], [2.1, 4.2], size=13, row_h=0.62)
table(s, Inches(7.1), BODY_Y, Inches(5.8), [
    ["Persona", "Goal"],
    ["Nusrat · learner", "Pass CSE211, spend ≤ ৳1,000"],
    ["Tanvir · tutor", "Earn about ৳10k a semester"],
    ["Dr. Farhana Rahman · advisor", "Stop the email ping-pong"],
], [2.9, 2.9], size=13, row_h=0.62)
tb(s, BODY_X, Inches(4.5), BODY_W, Inches(2.2), [
    "Task analysis (HTA) of “Subscribe to help for CSE211”: 4 sub-goals, 11 leaf steps (12 with demo)",
    "→ comparison data on result cards · one-screen checkout · demo email links straight to checkout",
], size=16, color=MUTED)

# 17 ── Wireframes
s = slide("Task 3 · Annotated wireframes: offering page and checkout", "Design")
half = (BODY_W - Inches(0.3)) // 2
image(s, os.path.join(T3, "wireframe-offering.png"), BODY_X, BODY_Y, half, BODY_H)
image(s, os.path.join(T3, "wireframe-checkout.png"), BODY_X + half + Inches(0.3), BODY_Y, half, BODY_H)

# 18 ── Login + Fitts
s = slide("Task 3 · Corrected login screen and Fitts's Law", "Design")
image(s, os.path.join(T3, "login-before-after.png"), BODY_X, BODY_Y, Inches(6.0), BODY_H)
table(s, Inches(6.9), BODY_Y, Inches(6.0), [
    ["MT = a + b·log₂(D/W + 1)", "A: small, top right", "B: sticky bottom"],
    ["Size", "90 × 36 px", "358 × 52 px"],
    ["D (from (195, 500))", "450.7 px", "304.0 px"],
    ["ID = log₂(D/W+1)", "3.76 bits", "2.78 bits"],
    ["MT (a = 0.10, b = 0.15)", "0.66 s", "0.52 s"],
], [2.4, 1.8, 1.8], size=13, row_h=0.5)
bullets(s, Inches(6.9), Inches(4.4), Inches(6.0), Inches(2.5), [
    "B is 22% faster|, in the thumb zone and above the 44–48 px minimum",
    "Login fixes|: real labels, an error under the field, one 342 × 52 button, no RESET, a sign-up path",
], size=15)

# 19 ── Next
s = slide("Next steps and questions for sir", "Plan")
bullets(s, BODY_X, BODY_Y, Inches(6.2), BODY_H, [
    "Next|",
    "  Task 4: function points (feeds COCOMO in Task 7), OO metrics, Halstead",
    "  Run the interviews and survey → confirm the refund rule and price",
    "  Agree team roles: PO, scrum master, QA lead",
    "  Build M1: sign-up → search → demo → mock payment",
], size=18)
bullets(s, Inches(7.0), BODY_Y, Inches(5.9), BODY_H, [
    "Questions for Masum sir|",
    "  Are the swimlane, sequence and class diagrams required?",
    "  Is a full SRS (ISO 29148) required, or is the PRD enough?",
    "  Do we use UniRide's figures (activity list, Halstead) or our own?",
    "  Which week is each task due?",
], size=18)

footer_numbers()
path = os.path.join(OUT, "proxy-courses-progress-tasks-1-3.pptx")
prs.save(path)
print(path)
